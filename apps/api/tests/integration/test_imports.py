"""Phase 8 acceptance (FACT): a realistic messy Excel sheet imports with clear
error reporting and correct opening balances (BACKLOG 8.1-8.4).

The workbook is synthetic (Q-35 default) but messy the way showroom sheets are:
a title above the table, Arabic headers spelled in different ways, Arabic-Indic
digits, currency words, day/month/year dates, an Excel date cell, a duplicate
chassis number, a missing make and an impossible date."""

import base64
import io
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import Connection, text

from app.core.errors import AppError
from app.domain.imports import EquityClearingIn, ImportCreateIn, ImportMappingIn
from app.services import finance, imports, partners
from tests.integration.conftest import NOUR, auth
from tests.integration.test_vehicles_sales import _NoCipher

D = Decimal


def _workbook(go_live: date, *, broken: bool) -> bytes:
    book = Workbook()
    cars = book.active
    assert cars is not None
    cars.title = "العربيات"
    cars.append(["كشف عربيات معرض الأمل — سبتمبر"])
    cars.append([])
    cars.append(
        [
            "م",
            "الماركه",
            "الموديل",
            "سنه الصنع",
            "رقم الشاسية",
            "سعر الشراء",
            "نقل",
            "صيانة",
            "تاريخ الشراء",
            "البائع",
            "باقي للبائع",
            "ملاحظات",
        ]
    )
    cars.append(
        [
            1,
            "Toyota",
            "Corolla",
            "٢٠٢٠",
            "JTDBR32E720011111",
            "٤٠٠٬٠٠٠",
            "EGP 2,500",
            "10,000 ج.م",
            "03/09/2026",
            None,
            None,
            "نظيفة",
        ]
    )
    cars.append(
        [
            2,
            "Hyundai",
            "Elantra",
            2019,
            "KMHD841CBKU222222",
            300000,
            None,
            "5000",
            datetime(2026, 9, 5),  # noqa: DTZ001 - Excel date cells carry no timezone
            "كريم محمود",
            "50,000",
            None,
        ]
    )
    if broken:
        cars.append([3, None, "Accent", 2018, "KMHC000000333333", "150000", None, None, "2026-09-10", None, None, None])
        cars.append([4, "Kia", "Rio", 2017, "JTDBR32E720011111", "abc", None, None, "31/02/2026", None, None, None])

    people = book.create_sheet("العملاء")
    people.append(["الإسم", "الموبايل", "العنوان"])
    people.append(["سامي فؤاد", "٠١٠٠٩٨٧٦٥٤٣", "الجيزة"])
    people.append(["هالة يوسف", "01112345678", None])

    owners = book.create_sheet("الشركاء")
    owners.append(["اسم الشريك", "النسبة", "رأس المال"])
    owners.append(["أحمد", "60%", "600000"])
    owners.append(["منى", "40", "400000"])

    plans = book.create_sheet("الأقساط")
    plans.append(["العميل", "الموبايل", "رقم العقد", "السيارة", "تاريخ الاستحقاق", "قيمة القسط", "المدفوع"])
    first_due = go_live + timedelta(days=10)
    plans.append(["مريم حسن", "01223344556", "C-17", "Kia Cerato", first_due.strftime("%d/%m/%Y"), "40000", "0"])
    plans.append(
        ["مريم حسن", "01223344556", "C-17", "Kia Cerato", (first_due + timedelta(days=30)).isoformat(), "40,000", None]
    )
    plans.append(
        ["مريم حسن", "01223344556", "C-17", "Kia Cerato", (first_due + timedelta(days=60)).isoformat(), "40000", "0"]
    )

    cash = book.create_sheet("الخزنة")
    cash.append(["الحساب", "النوع", "الرصيد"])
    cash.append(["الخزنة الرئيسية", "خزنة", "200,000"])
    cash.append(["بنك CIB", "بنك", "800000"])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _upload(conn: Connection, go_live: date, *, broken: bool) -> Any:
    content = base64.b64encode(_workbook(go_live, broken=broken)).decode()
    return imports.create(conn, ImportCreateIn(go_live_date=go_live, file_name="showroom.xlsx", content_base64=content))


def _balance(conn: Connection, key: str, **ids: uuid.UUID) -> Decimal:
    filters = "".join(f" and l.{column} = :{column}" for column in ids)
    return D(
        conn.execute(
            text(
                "select coalesce(sum(l.debit - l.credit), 0) from public.journal_lines l "  # noqa: S608 - test columns
                "join public.ledger_accounts a on a.id = l.ledger_account_id where a.system_key = :k" + filters
            ),
            {"k": key, **ids},
        ).scalar_one()
    )


def test_a_messy_sheet_reports_clear_errors_and_saves_nothing(fresh_tenant: tuple[Connection, uuid.UUID]) -> None:
    conn, _ = fresh_tenant
    go_live = finance.tenant_info(conn).today - timedelta(days=2)
    job = _upload(conn, go_live, broken=True)
    assert [s.kind for s in job.sheets] == ["VEHICLES", "CUSTOMERS", "PARTNERS", "INSTALLMENTS", "CASH"]
    cars = job.sheets[0]
    assert cars.header_row == 2
    assert cars.column_map == [
        None,
        "make",
        "model",
        "year",
        "vin",
        "purchase_price",
        "expense",
        "expense",
        "purchase_date",
        "seller_name",
        "owed_to_seller",
        "notes",
    ]
    checked = imports.validate(conn, job.id).sheets[0]
    assert checked.ok_rows == 2
    errors = {e.row_no: {(x.field, x.code) for x in e.errors} for e in checked.error_rows or []}
    assert errors[6] == {("make", "REQUIRED")}
    assert errors[7] == {("purchase_price", "BAD_NUMBER"), ("purchase_date", "BAD_DATE")}

    sheet = load_workbook(io.BytesIO(imports.errors_xlsx(conn, job.id, "ar")))["العربيات"]
    messages = [row[-1] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "make: مطلوب" in messages
    assert any("رقم غير صحيح" in (m or "") and "تاريخ غير صحيح" in (m or "") for m in messages)

    with pytest.raises(AppError) as refused:
        imports.commit(conn, job.id, _NoCipher())  # type: ignore[arg-type]
    assert refused.value.code == "IMPORT_HAS_ERRORS"
    assert conn.execute(text("select count(*) from public.vehicles")).scalar_one() == 0


def test_a_clean_sheet_posts_one_opening_entry(fresh_tenant: tuple[Connection, uuid.UUID]) -> None:
    conn, _ = fresh_tenant
    go_live = finance.tenant_info(conn).today - timedelta(days=2)
    job = _upload(conn, go_live, broken=False)
    validated = imports.validate(conn, job.id)
    assert all(not s.error_rows for s in validated.sheets)
    # Debits: cars 400,000 + 2,500 + 10,000 + 300,000 + 5,000; installments 120,000; cash 1,000,000.
    assert validated.opening_total == D("1837500.00")
    posted = imports.commit(conn, job.id, _NoCipher())  # type: ignore[arg-type]
    assert len(posted.journal_entries) == 1
    assert posted.document.result == {
        "vehicles": 2,
        "customers": 2,
        "partners": 2,
        "installment_plans": 1,
        "cash_accounts": 2,
    }

    entry = conn.execute(
        text("select is_opening, entry_date, source_type from public.journal_entries where id = :id"),
        {"id": posted.journal_entries[0].id},
    ).one()
    assert (entry.is_opening, entry.entry_date, entry.source_type) == (True, go_live, "OPENING_BALANCE")
    assert _balance(conn, "VEHICLE_INVENTORY") == D("717500.00")
    assert _balance(conn, "SELLER_PAYABLE") == D("-50000.00")
    assert _balance(conn, "INSTALLMENT_RECEIVABLE") == D("120000.00")
    assert _balance(conn, "PARTNER_CAPITAL") == D("-1000000.00")
    # 1,837,500 - 50,000 - 1,000,000 = 787,500 left on opening balance equity (Q-16).
    assert _balance(conn, "OPENING_BALANCE_EQUITY") == D("-787500.00")

    corolla = conn.execute(
        text("select id, status, stock_date from public.vehicles where vin_normalized = 'JTDBR32E720011111'")
    ).one()
    assert (corolla.status, corolla.stock_date) == ("AVAILABLE", date(2026, 9, 3))
    memos = set(
        conn.execute(text("select memo from public.journal_lines where vehicle_id = :v"), {"v": corolla.id}).scalars()
    )
    assert memos == {"سعر الشراء", "نقل", "صيانة"}
    plan = conn.execute(text("select opening_reference, installment_count from public.installment_plans")).one()
    assert (plan.opening_reference, plan.installment_count) == ("C-17", 3)
    shares = {s.partner_name_ar: s.percentage for s in partners.shares_on(conn, go_live)}
    assert shares == {"أحمد": D("60.0000"), "منى": D("40.0000")}

    # The confirmed mapping is remembered for the same layout.
    again = _upload(conn, go_live, broken=False)
    assert again.sheets[0].mapping_remembered

    with pytest.raises(AppError) as twice:
        imports.commit(conn, job.id, _NoCipher())  # type: ignore[arg-type]
    assert twice.value.code == "IMPORT_COMMITTED"

    # P-08: opening balance equity cleared to the partners by agreement.
    ahmed, mona = (s.partner_id for s in partners.shares_on(conn, go_live))
    clearing = EquityClearingIn.model_validate(
        {
            "clearing_date": go_live,
            "lines": [
                {"partner_id": ahmed, "account": "CAPITAL", "amount": "472500"},
                {"partner_id": mona, "account": "CURRENT", "amount": "315000.01"},
            ],
        }
    )
    with pytest.raises(AppError) as too_much:
        imports.clear_opening_equity(conn, clearing)
    assert too_much.value.code == "OPENING_EQUITY_EXCEEDED"
    clearing.lines[1].amount = D("315000")
    imports.clear_opening_equity(conn, clearing)
    assert _balance(conn, "OPENING_BALANCE_EQUITY") == 0


def test_partners_cannot_be_imported_twice(fresh_tenant: tuple[Connection, uuid.UUID]) -> None:
    conn, _ = fresh_tenant
    go_live = finance.tenant_info(conn).today - timedelta(days=2)
    job = _upload(conn, go_live, broken=False)
    imports.validate(conn, job.id)
    imports.commit(conn, job.id, _NoCipher())  # type: ignore[arg-type]
    second = _upload(conn, go_live, broken=False)
    keep_partners_only = ImportMappingIn.model_validate(
        {
            "sheets": [
                {"index": s.index, "column_map": s.column_map, "skip": s.kind != "PARTNERS"} for s in second.sheets
            ]
        }
    )
    imports.set_mapping(conn, second.id, keep_partners_only)
    partners_sheet = next(s for s in imports.validate(conn, second.id).sheets if s.kind == "PARTNERS")
    assert {e.code for r in partners_sheet.error_rows or [] for e in r.errors} == {"PARTNERS_EXIST"}


def test_rows_from_the_onboarding_wizard(fresh_tenant: tuple[Connection, uuid.UUID]) -> None:
    conn, _ = fresh_tenant
    go_live = finance.tenant_info(conn).today
    job = imports.create(
        conn,
        ImportCreateIn.model_validate(
            {
                "go_live_date": go_live,
                "sheets": [
                    {
                        "name": "partners",
                        "kind": "PARTNERS",
                        "headers": ["name", "percentage", "capital"],
                        "rows": [["خالد", "100", "250000"]],
                    },
                    {
                        "name": "cash",
                        "kind": "CASH",
                        "headers": ["account_name", "account_type", "balance"],
                        "rows": [["الخزنة", "CASH_BOX", "250000"]],
                    },
                ],
            }
        ),
    )
    imports.validate(conn, job.id)
    imports.commit(conn, job.id, _NoCipher())  # type: ignore[arg-type]
    assert _balance(conn, "OPENING_BALANCE_EQUITY") == 0
    assert _balance(conn, "PARTNER_CAPITAL") == D("-250000.00")


def test_only_owners_import(client: Any) -> None:
    sales = client.get("/api/v1/imports/template", headers=auth("sales@nour.example", NOUR))
    assert sales.status_code == 403
    owner = client.get("/api/v1/imports/template", headers=auth("owner@nour.example", NOUR))
    assert owner.status_code == 200
    assert load_workbook(io.BytesIO(owner.content)).sheetnames[0] == "السيارات"
