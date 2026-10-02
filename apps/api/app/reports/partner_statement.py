"""Partner statement export (كشف حساب شريك, SPEC §4.2): Excel and PDF, Arabic-first."""

from decimal import Decimal
from html import escape
from io import BytesIO
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.money import format_money
from app.domain.partners import PartnerPosition, PartnerStatementOut
from app.reports.cash_book import FONTS, html_to_pdf, sheet_title

Language = Literal["ar", "en"]

_L: dict[str, dict[Language, str]] = {
    "title": {"ar": "كشف حساب شريك", "en": "Partner statement"},
    "partner": {"ar": "الشريك", "en": "Partner"},
    "period": {"ar": "الفترة", "en": "Period"},
    "opening": {"ar": "الرصيد الافتتاحي", "en": "Opening"},
    "closing": {"ar": "الرصيد الختامي", "en": "Closing"},
    "capital": {"ar": "رأس المال", "en": "Capital"},
    "current": {"ar": "الحساب الجاري", "en": "Current account"},
    "loans_to": {"ar": "سلف على الشريك", "en": "Loans to partner"},
    "loans_from": {"ar": "قروض من الشريك", "en": "Loans from partner"},
    "net": {"ar": "صافي مركز الشريك", "en": "Net position"},
    "date": {"ar": "التاريخ", "en": "Date"},
    "entry": {"ar": "رقم القيد", "en": "Entry #"},
    "description": {"ar": "البيان", "en": "Description"},
    "bucket": {"ar": "البند", "en": "Bucket"},
    "in": {"ar": "له", "en": "In his favour"},
    "out": {"ar": "عليه", "en": "Against"},
    "running": {"ar": "الرصيد", "en": "Balance"},
    "empty": {"ar": "لا توجد حركات في هذه الفترة", "en": "No movements in this period"},
    "CAPITAL": {"ar": "رأس المال", "en": "Capital"},
    "CURRENT": {"ar": "جاري", "en": "Current"},
    "LOAN_TO": {"ar": "سلفة", "en": "Loan to partner"},
    "LOAN_FROM": {"ar": "قرض للمعرض", "en": "Loan from partner"},
}


def _l(key: str, language: Language) -> str:
    return _L[key][language]


def _name(statement: PartnerStatementOut, language: Language) -> str:
    partner = statement.partner
    return (partner.name_en or partner.name_ar) if language == "en" else partner.name_ar


def filename(statement: PartnerStatementOut, extension: str) -> str:
    return f"partner-statement-{statement.date_from}-{statement.date_to}.{extension}"


def _position_rows(position: PartnerPosition) -> list[tuple[str, Decimal]]:
    return [
        ("capital", position.capital),
        ("current", position.current),
        ("loans_to", position.loans_to_partner),
        ("loans_from", position.loans_from_partner),
        ("net", position.net),
    ]


def render_xlsx(statement: PartnerStatementOut, language: Language) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("workbook has no sheet")
    sheet.title = sheet_title(_l("title", language))
    sheet.sheet_view.rightToLeft = language == "ar"
    bold = Font(bold=True)

    sheet.append([_l("title", language)])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([_l("partner", language), _name(statement, language)])
    sheet.append([_l("period", language), f"{statement.date_from} → {statement.date_to}"])
    sheet.append([])
    sheet.append(["", _l("opening", language), _l("closing", language)])
    for cell in sheet[sheet.max_row]:
        cell.font = bold
    for (key, opening), (_, closing) in zip(
        _position_rows(statement.opening), _position_rows(statement.closing), strict=True
    ):
        sheet.append([_l(key, language), float(opening), float(closing)])
    sheet.append([])

    sheet.append([_l(k, language) for k in ("date", "entry", "description", "bucket", "in", "out", "running")])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="E2E8F0")
    for line in statement.lines:
        sheet.append(
            [
                line.entry_date,
                line.entry_no,
                line.description,
                _l(line.bucket, language),
                float(line.amount_in) or None,
                float(line.amount_out) or None,
                float(line.running_net),
            ]
        )

    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
            elif cell.column == 1 and (cell.row or 0) > header_row and not isinstance(cell.value, str | None):
                cell.number_format = "yyyy-mm-dd"
    for column, width in zip(range(1, 8), (16, 16, 40, 18, 16, 16, 18), strict=True):
        sheet.column_dimensions[get_column_letter(column)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_html(statement: PartnerStatementOut, language: Language, showroom_name: str) -> str:
    direction = "rtl" if language == "ar" else "ltr"

    def money(value: Decimal) -> str:
        return escape(format_money(value, statement.currency_code, language))

    positions = "".join(
        f"<tr><td>{_l(key, language)}</td><td class='num'>{money(opening)}</td><td class='num'>{money(closing)}</td></tr>"
        for (key, opening), (_, closing) in zip(
            _position_rows(statement.opening), _position_rows(statement.closing), strict=True
        )
    )
    rows = (
        "".join(
            "<tr>"
            f"<td>{line.entry_date}</td><td>{line.entry_no}</td><td>{escape(line.description)}</td>"
            f"<td>{_l(line.bucket, language)}</td>"
            f"<td class='num'>{money(line.amount_in) if line.amount_in else ''}</td>"
            f"<td class='num'>{money(line.amount_out) if line.amount_out else ''}</td>"
            f"<td class='num'>{money(line.running_net)}</td>"
            "</tr>"
            for line in statement.lines
        )
        or f"<tr><td colspan='7' class='empty'>{_l('empty', language)}</td></tr>"
    )
    fonts = FONTS.as_uri()
    return f"""<!doctype html>
<html lang="{language}" dir="{direction}"><head><meta charset="utf-8"><style>
  @font-face {{ font-family: Plex; src: url('{fonts}/ibm-plex-sans-arabic-arabic-400-normal.woff'); font-weight: 400; }}
  @font-face {{ font-family: Plex; src: url('{fonts}/ibm-plex-sans-arabic-arabic-700-normal.woff'); font-weight: 700; }}
  @font-face {{ font-family: PlexLatin; src: url('{fonts}/ibm-plex-sans-arabic-latin-400-normal.woff'); font-weight: 400; }}
  @font-face {{ font-family: PlexLatin; src: url('{fonts}/ibm-plex-sans-arabic-latin-700-normal.woff'); font-weight: 700; }}
  @page {{ size: A4 landscape; margin: 14mm; @bottom-center {{ content: counter(page) " / " counter(pages); font-size: 8pt; }} }}
  body {{ font-family: PlexLatin, Plex, sans-serif; font-size: 9pt; color: #0f172a; }}
  h1 {{ font-size: 15pt; margin: 0 0 3mm; }}
  table {{ width: 100%; border-collapse: collapse; margin-bottom: 5mm; }}
  th {{ background: #e2e8f0; text-align: start; padding: 2mm; }}
  td {{ border-bottom: 0.3pt solid #cbd5e1; padding: 1.5mm 2mm; }}
  .num {{ text-align: end; white-space: nowrap; }}
  .positions {{ width: 60%; }}
  .empty {{ text-align: center; color: #475569; padding: 6mm; }}
</style></head><body>
  <h1>{escape(showroom_name)} — {_l("title", language)}</h1>
  <p><b>{_l("partner", language)}:</b> {escape(_name(statement, language))} &nbsp;
     <b>{_l("period", language)}:</b> {statement.date_from} — {statement.date_to}</p>
  <table class="positions"><thead><tr><th></th><th class="num">{_l("opening", language)}</th>
    <th class="num">{_l("closing", language)}</th></tr></thead><tbody>{positions}</tbody></table>
  <table><thead><tr>
    <th>{_l("date", language)}</th><th>{_l("entry", language)}</th><th>{_l("description", language)}</th>
    <th>{_l("bucket", language)}</th><th class="num">{_l("in", language)}</th><th class="num">{_l("out", language)}</th>
    <th class="num">{_l("running", language)}</th></tr></thead><tbody>{rows}</tbody></table>
</body></html>"""


def render_pdf(statement: PartnerStatementOut, language: Language, showroom_name: str) -> bytes:
    return html_to_pdf(render_html(statement, language, showroom_name))
