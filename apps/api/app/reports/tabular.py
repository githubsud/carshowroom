"""Any ``ReportTable`` as PDF (WeasyPrint) or Excel (openpyxl), Arabic-first.

Text cells may carry an Arabic/English pair as "عربي|English"; the exporter
picks the side for the language. Money is written to Excel as numbers so
owners can keep calculating.
"""

from decimal import Decimal
from html import escape
from io import BytesIO
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.money import format_money
from app.domain.reports import ReportTable
from app.reports.cash_book import FONTS, html_to_pdf, sheet_title

Language = Literal["ar", "en"]


def pick(value: str | int | None, language: Language) -> str:
    if value is None:
        return ""
    text = str(value)
    if "|" in text:
        ar, _, en = text.partition("|")
        return ar if language == "ar" else en
    return text


def _title(table: ReportTable, language: Language) -> str:
    return table.title_ar if language == "ar" else table.title_en


def _period(table: ReportTable, language: Language) -> str:
    return (table.period_ar if language == "ar" else table.period_en) or ""


def filename(table: ReportTable, extension: str) -> str:
    return f"{table.name}.{extension}"


def render_xlsx(table: ReportTable, language: Language) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # a new workbook always has one sheet
        raise RuntimeError("workbook has no sheet")
    sheet.title = sheet_title(_title(table, language))
    sheet.sheet_view.rightToLeft = language == "ar"
    sheet.append([_title(table, language)])
    sheet["A1"].font = Font(bold=True, size=14)
    if _period(table, language):
        sheet.append([_period(table, language)])
    for label_ar, label_en, value, kind in table.figures:
        sheet.append([label_ar if language == "ar" else label_en, float(value) if kind == "money" else value])
    sheet.append([])
    sheet.append([c.label_ar if language == "ar" else c.label_en for c in table.columns])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2E8F0")
    for row in table.rows:
        values: list[object] = []
        for column in table.columns:
            raw = row.cells.get(column.key)
            if raw is None or raw == "":
                values.append(None)
            elif column.kind in ("money", "percent"):
                values.append(float(Decimal(str(raw))))
            elif column.kind == "number":
                values.append(int(raw))
            else:
                values.append(pick(raw, language))
        sheet.append(values)
        if row.style in ("total", "section"):
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
    for index, column in enumerate(table.columns, start=1):
        width = 18 if column.kind in ("money", "date") else 12 if column.kind in ("number", "percent") else 32
        sheet.column_dimensions[get_column_letter(index)].width = width
        if column.kind == "money":
            for (cell,) in sheet.iter_rows(min_row=header_row + 1, min_col=index, max_col=index):
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_html(table: ReportTable, language: Language, showroom_name: str) -> str:
    def cell(value: str | int | None, kind: str) -> str:
        if value is None or value == "":
            return "<td></td>"
        if kind == "money":
            return f"<td class='num'>{escape(format_money(Decimal(str(value)), table.currency_code, language))}</td>"
        if kind in ("number", "percent"):
            return f"<td class='num'>{escape(str(value))}</td>"
        return f"<td>{escape(pick(value, language))}</td>"

    head = "".join(
        f"<th class='{'num' if c.kind in ('money', 'number', 'percent') else ''}'>"
        f"{escape(c.label_ar if language == 'ar' else c.label_en)}</th>"
        for c in table.columns
    )
    body = "".join(
        f"<tr class='{row.style or ''}'>" + "".join(cell(row.cells.get(c.key), c.kind) for c in table.columns) + "</tr>"
        for row in table.rows
    )
    figures = "".join(
        f"<div><span>{escape(ar if language == 'ar' else en)}</span> <b>"
        f"{escape(format_money(Decimal(value), table.currency_code, language) if kind == 'money' else value)}</b></div>"
        for ar, en, value, kind in table.figures
    )
    fonts = FONTS.as_uri()
    landscape = " landscape" if len(table.columns) > 6 else ""
    return f"""<!doctype html>
<html lang="{language}" dir="{"rtl" if language == "ar" else "ltr"}"><head><meta charset="utf-8"><style>
  @font-face {{ font-family: Plex; src: url('{fonts}/ibm-plex-sans-arabic-arabic-400-normal.woff'); font-weight: 400; }}
  @font-face {{ font-family: Plex; src: url('{fonts}/ibm-plex-sans-arabic-arabic-700-normal.woff'); font-weight: 700; }}
  @font-face {{ font-family: PlexLatin; src: url('{fonts}/ibm-plex-sans-arabic-latin-400-normal.woff'); font-weight: 400; }}
  @font-face {{ font-family: PlexLatin; src: url('{fonts}/ibm-plex-sans-arabic-latin-700-normal.woff'); font-weight: 700; }}
  @page {{ size: A4{landscape}; margin: 14mm; @bottom-center {{ content: counter(page) " / " counter(pages); font-size: 8pt; }} }}
  body {{ font-family: PlexLatin, Plex, sans-serif; font-size: 9pt; color: #0f172a; }}
  h1 {{ font-size: 15pt; margin: 0 0 2mm; }}
  .period {{ color: #475569; margin: 0 0 4mm; }}
  .figures {{ display: flex; flex-wrap: wrap; gap: 8mm; margin-bottom: 4mm; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th {{ background: #e2e8f0; text-align: start; padding: 1.8mm 2mm; }}
  td {{ border-bottom: 0.3pt solid #cbd5e1; padding: 1.4mm 2mm; vertical-align: top; }}
  .num {{ text-align: end; white-space: nowrap; }}
  tr.total td {{ font-weight: 700; border-top: 0.6pt solid #0f172a; }}
  tr.section td {{ font-weight: 700; background: #f8fafc; }}
  tr.muted td {{ color: #64748b; }}
</style></head><body>
  <h1>{escape(showroom_name)} — {escape(_title(table, language))}</h1>
  <p class="period">{escape(_period(table, language))}</p>
  <div class="figures">{figures}</div>
  <table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>
</body></html>"""


def render_pdf(table: ReportTable, language: Language, showroom_name: str) -> bytes:
    return html_to_pdf(render_html(table, language, showroom_name))
