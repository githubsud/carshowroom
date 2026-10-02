"""Cash book export: Excel (openpyxl) and PDF (WeasyPrint), Arabic-first (SPEC §4.12).

The numbers come from app.services.finance.cash_book; this module only lays
them out. Amounts are written to Excel as numbers with 2-decimal formatting so
owners can keep calculating in Excel.
"""

from decimal import Decimal
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.errors import AppError
from app.domain.finance import CashBookOut
from app.domain.money import format_money

Language = Literal["ar", "en"]
FONTS = Path(__file__).parent / "fonts"

_LABELS: dict[str, dict[Language, str]] = {
    "title": {"ar": "دفتر الخزنة / البنك", "en": "Cash / bank book"},
    "account": {"ar": "الحساب", "en": "Account"},
    "period": {"ar": "الفترة", "en": "Period"},
    "opening": {"ar": "الرصيد الافتتاحي", "en": "Opening balance"},
    "closing": {"ar": "الرصيد الختامي", "en": "Closing balance"},
    "total_in": {"ar": "إجمالي الوارد", "en": "Total in"},
    "total_out": {"ar": "إجمالي المنصرف", "en": "Total out"},
    "date": {"ar": "التاريخ", "en": "Date"},
    "entry": {"ar": "رقم القيد", "en": "Entry #"},
    "description": {"ar": "البيان", "en": "Description"},
    "against": {"ar": "الطرف المقابل", "en": "Against"},
    "in": {"ar": "وارد", "en": "In"},
    "out": {"ar": "منصرف", "en": "Out"},
    "balance": {"ar": "الرصيد", "en": "Balance"},
    "reversal": {"ar": "قيد عكسي", "en": "reversal"},
    "reversed": {"ar": "معكوس", "en": "reversed"},
    "empty": {"ar": "لا توجد حركات في هذه الفترة", "en": "No movements in this period"},
}


def _label(key: str, language: Language) -> str:
    return _LABELS[key][language]


def _account_name(book: CashBookOut, language: Language) -> str:
    account = book.cash_account
    return account.name_en or account.name_ar if language == "en" else account.name_ar


def sheet_title(title: str) -> str:
    """Excel forbids the characters [ ] : * ? / and backslash in sheet names, and allows 31 characters."""
    for forbidden in "[]:*?/\\":
        title = title.replace(forbidden, "-")
    return title[:31]


def filename(book: CashBookOut, extension: str) -> str:
    return f"cash-book-{book.cash_account.ledger_account_code}-{book.date_from}-{book.date_to}.{extension}"


# --- Excel -------------------------------------------------------------------------------------


def render_xlsx(book: CashBookOut, language: Language) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # a new workbook always has one sheet
        raise RuntimeError("workbook has no sheet")
    sheet.title = sheet_title(_label("title", language))
    sheet.sheet_view.rightToLeft = language == "ar"

    bold = Font(bold=True)
    sheet.append([_label("title", language)])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([_label("account", language), _account_name(book, language)])
    sheet.append([_label("period", language), f"{book.date_from} → {book.date_to}"])
    sheet.append([_label("opening", language), float(book.opening_balance)])
    sheet.append([])

    header = ["date", "entry", "description", "against", "in", "out", "balance"]
    sheet.append([_label(key, language) for key in header])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="E2E8F0")

    for movement in book.movements:
        description = movement.description
        if movement.is_reversal:
            description += f" ({_label('reversal', language)})"
        elif movement.reversed:
            description += f" ({_label('reversed', language)})"
        sheet.append(
            [
                movement.entry_date,
                movement.entry_no,
                description,
                movement.counterpart_ar if language == "ar" else movement.counterpart_en,
                float(movement.amount_in) or None,
                float(movement.amount_out) or None,
                float(movement.balance),
            ]
        )

    sheet.append([])
    for key, value in (
        ("total_in", book.total_in),
        ("total_out", book.total_out),
        ("closing", book.closing_balance),
    ):
        sheet.append([_label(key, language), float(value)])
        sheet.cell(row=sheet.max_row, column=1).font = bold

    money_format = "#,##0.00;[Red]-#,##0.00"
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = money_format
            elif cell.column == 1 and (cell.row or 0) > header_row and not isinstance(cell.value, str | None):
                cell.number_format = "yyyy-mm-dd"
    for column, width in zip(range(1, 8), (14, 10, 40, 30, 16, 16, 18), strict=True):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for cell in sheet["C"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- PDF -----------------------------------------------------------------------------------------


def render_html(book: CashBookOut, language: Language, showroom_name: str) -> str:
    direction = "rtl" if language == "ar" else "ltr"

    def money(value: Decimal) -> str:
        return format_money(value, book.currency_code, language)

    def amount_cell(value: Decimal) -> str:
        return f"<td class='num'>{escape(money(value))}</td>" if value else "<td class='num'></td>"

    rows = []
    for m in book.movements:
        tag = ""
        if m.is_reversal:
            tag = f" <span class='tag'>{_label('reversal', language)}</span>"
        elif m.reversed:
            tag = f" <span class='tag'>{_label('reversed', language)}</span>"
        rows.append(
            "<tr>"
            f"<td>{m.entry_date}</td><td>{m.entry_no}</td>"
            f"<td>{escape(m.description)}{tag}</td>"
            f"<td>{escape(m.counterpart_ar if language == 'ar' else m.counterpart_en)}</td>"
            f"{amount_cell(m.amount_in)}{amount_cell(m.amount_out)}"
            f"<td class='num'>{escape(money(m.balance))}</td>"
            "</tr>"
        )
    body_rows = "".join(rows) or f"<tr><td colspan='7' class='empty'>{_label('empty', language)}</td></tr>"
    fonts = FONTS.as_uri()
    return f"""<!doctype html>
<html lang="{language}" dir="{direction}">
<head><meta charset="utf-8">
<style>
  @font-face {{ font-family: Plex; src: url('{fonts}/ibm-plex-sans-arabic-arabic-400-normal.woff'); font-weight: 400; }}
  @font-face {{ font-family: Plex; src: url('{fonts}/ibm-plex-sans-arabic-arabic-700-normal.woff'); font-weight: 700; }}
  @font-face {{ font-family: PlexLatin; src: url('{fonts}/ibm-plex-sans-arabic-latin-400-normal.woff'); font-weight: 400; }}
  @font-face {{ font-family: PlexLatin; src: url('{fonts}/ibm-plex-sans-arabic-latin-700-normal.woff'); font-weight: 700; }}
  @page {{ size: A4 landscape; margin: 14mm; @bottom-center {{ content: counter(page) " / " counter(pages); font-size: 8pt; }} }}
  body {{ font-family: PlexLatin, Plex, sans-serif; font-size: 9pt; color: #0f172a; }}
  h1 {{ font-size: 15pt; margin: 0 0 4mm; }}
  .meta {{ display: flex; gap: 10mm; margin-bottom: 4mm; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th {{ background: #e2e8f0; text-align: start; padding: 2mm; }}
  td {{ border-bottom: 0.3pt solid #cbd5e1; padding: 1.5mm 2mm; vertical-align: top; }}
  .num {{ text-align: end; white-space: nowrap; }}
  .totals td {{ font-weight: 700; border-bottom: none; }}
  .tag {{ font-size: 7pt; color: #b91c1c; }}
  .empty {{ text-align: center; color: #475569; padding: 6mm; }}
</style></head>
<body>
  <h1>{escape(showroom_name)} — {_label("title", language)}</h1>
  <div class="meta">
    <div><b>{_label("account", language)}:</b> {escape(_account_name(book, language))}
         ({book.cash_account.ledger_account_code})</div>
    <div><b>{_label("period", language)}:</b> {book.date_from} — {book.date_to}</div>
    <div><b>{_label("opening", language)}:</b> {escape(money(book.opening_balance))}</div>
  </div>
  <table>
    <thead><tr>
      <th>{_label("date", language)}</th><th>{_label("entry", language)}</th>
      <th>{_label("description", language)}</th><th>{_label("against", language)}</th>
      <th class="num">{_label("in", language)}</th><th class="num">{_label("out", language)}</th>
      <th class="num">{_label("balance", language)}</th>
    </tr></thead>
    <tbody>{body_rows}</tbody>
    <tfoot class="totals"><tr>
      <td colspan="4">{_label("closing", language)}</td>
      <td class="num">{escape(money(book.total_in))}</td><td class="num">{escape(money(book.total_out))}</td>
      <td class="num">{escape(money(book.closing_balance))}</td>
    </tr></tfoot>
  </table>
</body></html>"""


def render_pdf(book: CashBookOut, language: Language, showroom_name: str) -> bytes:
    try:
        from weasyprint import HTML
    except OSError as exc:
        raise AppError(
            "PDF_UNAVAILABLE",
            "PDF rendering is not available on this server (missing Pango libraries)",
            status_code=503,
        ) from exc
    pdf: bytes = HTML(string=render_html(book, language, showroom_name)).write_pdf()
    return pdf
