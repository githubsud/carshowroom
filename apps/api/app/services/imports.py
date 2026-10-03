"""Excel import and opening balances (SPEC §4.13, BACKLOG 8.1-8.4).

1. Upload (their own .xlsx/.csv, our template, or rows from the onboarding
   wizard): every sheet is read as text, the header row found, the kind and
   the column mapping suggested — a mapping saved for the same header layout
   wins (D-109).
2. Mapping: the user confirms or changes it; it is remembered per tenant.
3. Validation: every row is parsed (Arabic-Indic digits, currency symbols,
   d/m/y dates, Excel serials, phones) and checked (required fields,
   duplicates in the file and in the database, shares totalling 100%);
   nothing is saved, and the errors can be downloaded as a sheet.
4. Commit: all or nothing, in one transaction — vehicles (with their cost
   breakdown), customers, partners and their shares, open installments, cash
   and bank accounts — and ONE opening entry (rule 25) at the go-live date,
   balanced by opening balance equity (3900, Q-16 / P-08).
"""

import base64
import csv
import io
import json
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from uuid import UUID, uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import Connection, text

from app.core.crypto import FieldCipher
from app.core.errors import AppError, not_found
from app.core.phone import normalize_phone
from app.domain.finance import CashAccountIn, EntryRef, PostingResult
from app.domain.imports import (
    FIELDS,
    EquityClearingIn,
    FieldSpec,
    ImportCreateIn,
    ImportJobOut,
    ImportKind,
    ImportMappingIn,
    OpeningEquityOut,
    ParseError,
    RowError,
    RowResult,
    SheetOut,
    detect_header_row,
    detect_kind,
    normalize_header,
    parse_date,
    parse_number,
    suggest_mapping,
    to_western_digits,
)
from app.domain.ledger import ZERO, Account, Line
from app.domain.partners import PartnerIn, ShareChangeIn
from app.domain.vehicles import CustomerIn, VehicleIn
from app.services import customers, finance, partners, vehicles
from app.services.posting import engine, rules

MAX_ROWS = 5000
_HUNDRED = Decimal(100)


# =====================================================================================================
# Reading files
# =====================================================================================================


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    text_value = str(value).strip()
    return text_value or None


def _read_file(file_name: str, content: bytes) -> list[tuple[str, list[list[str | None]]]]:
    name = file_name.lower()
    if name.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1256"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise AppError("UPLOAD_INVALID", "The file could not be read", status_code=422)
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t") if decoded.strip() else csv.excel
        rows = [[_cell_text(c) for c in row] for row in csv.reader(io.StringIO(decoded), dialect)]
        return [(file_name.rsplit(".", 1)[0], rows)]
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds of errors for a bad file
        raise AppError("UPLOAD_INVALID", "Upload an .xlsx or .csv file", status_code=422) from exc
    sheets = []
    for sheet in workbook.worksheets:
        rows = [[_cell_text(c) for c in row] for row in sheet.iter_rows(values_only=True)]
        if any(any(c is not None for c in row) for row in rows):
            sheets.append((sheet.title, rows))
    return sheets


def _trim(row: list[str | None], width: int) -> list[str | None]:
    padded = list(row[:width]) + [None] * max(0, width - len(row))
    return padded


def _signature(headers: list[str]) -> str:
    return "|".join(normalize_header(h) for h in headers)


def _saved_mapping(conn: Connection, kind: ImportKind, headers: list[str]) -> list[str | None] | None:
    value = conn.execute(
        text("select column_map from public.import_mappings where kind = :kind and signature = :sig"),
        {"kind": kind, "sig": _signature(headers)},
    ).scalar_one_or_none()
    if value is not None and len(value) == len(headers):
        return list(value)
    return None


def _build_sheet(
    conn: Connection, name: str, raw: list[list[str | None]], *, kind: ImportKind | None, keys: bool
) -> dict[str, Any]:
    header_row = 0 if keys else detect_header_row([list(r) for r in raw])
    header_cells = raw[header_row] if raw else []
    width = max((i + 1 for i, c in enumerate(header_cells) if c not in (None, "")), default=0)
    headers = [str(c or f"#{i + 1}") for i, c in enumerate(header_cells[:width])]
    rows = [_trim(r, width) for r in raw[header_row + 1 :]]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if len(rows) > MAX_ROWS:
        raise AppError("IMPORT_TOO_LARGE", f"At most {MAX_ROWS} rows per sheet", status_code=422)
    sheet_kind = kind or detect_kind(headers, name)
    saved = None if keys else _saved_mapping(conn, sheet_kind, headers)
    if keys:
        valid = {spec.key for spec in FIELDS[sheet_kind]}
        column_map: list[str | None] = [h if h in valid else None for h in headers]
    else:
        column_map = saved or suggest_mapping(headers, sheet_kind)
    return {
        "name": name,
        "kind": sheet_kind,
        "skip": False,
        "header_row": header_row,
        "headers": headers,
        "rows": rows,
        "column_map": column_map,
        "remembered": saved is not None,
    }


# =====================================================================================================
# Jobs
# =====================================================================================================


def create(conn: Connection, payload: ImportCreateIn) -> ImportJobOut:
    info = finance.tenant_info(conn)
    if payload.go_live_date > info.today:
        raise AppError("DATE_IN_FUTURE", "The go-live date cannot be in the future", status_code=422)
    sheets: list[dict[str, Any]] = []
    if payload.content_base64:
        try:
            content = base64.b64decode(payload.content_base64, validate=True)
        except ValueError as exc:
            raise AppError("UPLOAD_INVALID", "The file could not be read", status_code=422) from exc
        for name, raw in _read_file(payload.file_name or "import.xlsx", content):
            sheets.append(_build_sheet(conn, name, raw, kind=None, keys=False))
    for given in payload.sheets:
        given_rows: list[list[str | None]] = [list(given.headers), *given.rows]
        sheets.append(_build_sheet(conn, given.name, given_rows, kind=given.kind, keys=True))
    if not sheets:
        raise AppError("UPLOAD_MISSING", "Upload a file or give the rows", status_code=422)
    job_id = conn.execute(
        text(
            """
            insert into public.import_jobs (tenant_id, file_name, go_live_date, sheets)
            values (private.current_tenant_id(), :file_name, :go_live, cast(:sheets as jsonb))
            returning id
            """
        ),
        {
            "file_name": payload.file_name,
            "go_live": payload.go_live_date,
            "sheets": json.dumps(sheets, ensure_ascii=False),
        },
    ).scalar_one()
    return get(conn, job_id)


_JOB = text(
    "select j.*, e.entry_no from public.import_jobs j "
    "left join public.journal_entries e on e.id = j.journal_entry_id where j.id = :id"
)


def _row(conn: Connection, job_id: UUID, *, lock: bool = False) -> Any:
    if lock:
        conn.execute(text("select 1 from public.import_jobs where id = :id for update"), {"id": job_id})
    row = conn.execute(_JOB, {"id": job_id}).mappings().first()
    if row is None:
        raise not_found("import")
    return row


def get(conn: Connection, job_id: UUID) -> ImportJobOut:
    row = _row(conn, job_id)
    validation = row["validation"] or {}
    by_sheet = {item["index"]: item for item in validation.get("sheets", [])}
    sheets = []
    for index, sheet in enumerate(row["sheets"]):
        checked = by_sheet.get(index)
        sheets.append(
            SheetOut(
                index=index,
                name=sheet["name"],
                kind=sheet["kind"],
                skip=sheet.get("skip", False),
                header_row=sheet["header_row"],
                headers=sheet["headers"],
                sample_rows=sheet["rows"][:5],
                row_count=len(sheet["rows"]),
                column_map=sheet["column_map"],
                mapping_remembered=sheet.get("remembered", False),
                ok_rows=checked["ok_rows"] if checked else None,
                error_rows=[RowResult.model_validate(e) for e in checked["errors"]] if checked else None,
            )
        )
    return ImportJobOut(
        id=row["id"],
        file_name=row["file_name"],
        go_live_date=row["go_live_date"],
        status=row["status"],
        sheets=sheets,
        fields={kind: FIELDS[kind] for kind in FIELDS},
        opening_total=Decimal(validation["opening_total"]) if validation.get("opening_total") is not None else None,
        result=row["result"],
        entry_no=row["entry_no"],
        created_at=row["created_at"],
    )


def list_jobs(conn: Connection) -> list[ImportJobOut]:
    ids = conn.execute(text("select id from public.import_jobs order by created_at desc limit 20")).scalars()
    return [get(conn, job_id) for job_id in ids]


def _remember(conn: Connection, sheet: dict[str, Any]) -> None:
    """The same header layout gets the same mapping next time (D-109)."""
    conn.execute(
        text(
            """
            insert into public.import_mappings (tenant_id, kind, signature, column_map)
            values (private.current_tenant_id(), :kind, :sig, cast(:map as jsonb))
            on conflict (tenant_id, kind, signature) do update set column_map = excluded.column_map
            """
        ),
        {"kind": sheet["kind"], "sig": _signature(sheet["headers"]), "map": json.dumps(sheet["column_map"])},
    )


def set_mapping(conn: Connection, job_id: UUID, payload: ImportMappingIn) -> ImportJobOut:
    row = _row(conn, job_id, lock=True)
    if row["status"] == "COMMITTED":
        raise AppError("IMPORT_COMMITTED", "This import has been committed", status_code=409)
    sheets = row["sheets"]
    for item in payload.sheets:
        if not 0 <= item.index < len(sheets):
            raise AppError("VALIDATION_ERROR", "Unknown sheet", status_code=422)
        sheet = sheets[item.index]
        if len(item.column_map) != len(sheet["headers"]):
            raise AppError("VALIDATION_ERROR", "The mapping must cover every column", status_code=422)
        kind: ImportKind = item.kind or sheet["kind"]
        valid = {spec.key for spec in FIELDS[kind]}
        if any(key is not None and key not in valid for key in item.column_map):
            raise AppError("VALIDATION_ERROR", "Unknown field in the mapping", status_code=422)
        sheet.update({"kind": kind, "column_map": item.column_map, "skip": item.skip})
        if item.remember and not item.skip:
            _remember(conn, sheet)
    conn.execute(
        text(
            "update public.import_jobs set sheets = cast(:sheets as jsonb), status = 'UPLOADED', validation = null "
            "where id = :id"
        ),
        {"sheets": json.dumps(sheets, ensure_ascii=False), "id": job_id},
    )
    return get(conn, job_id)


# =====================================================================================================
# Validation
# =====================================================================================================


@dataclass
class _Record:
    row_no: int
    values: dict[str, Any]
    expenses: list[tuple[str, Decimal]] = field(default_factory=list)


@dataclass
class _Checked:
    index: int
    kind: ImportKind
    records: list[_Record]
    errors: list[RowResult]


def _rows(sheet: dict[str, Any]) -> Iterator[tuple[int, dict[str, list[tuple[str, str]]]]]:
    """Each data row as {field: [(header, raw text), ...]} (row_no as in the sheet)."""
    first = sheet["header_row"] + 2  # 1-based, after the header
    for offset, row in enumerate(sheet["rows"]):
        cells: dict[str, list[tuple[str, str]]] = {}
        for header, key, raw in zip(sheet["headers"], sheet["column_map"], row, strict=False):
            if key is not None and raw not in (None, ""):
                cells.setdefault(key, []).append((header, str(raw)))
        yield first + offset, cells


def _parse(spec: FieldSpec, raw: str, country: str) -> Any:
    if spec.type == "money":
        return parse_number(raw)
    if spec.type == "percent":
        value = parse_number(raw.replace("%", ""))
        if not 0 < value <= 100:
            raise ParseError("BAD_PERCENT")
        return value
    if spec.type == "int":
        number = parse_number(raw)
        if number < 0 or number != number.to_integral_value():
            raise ParseError("BAD_NUMBER")
        return int(number)
    if spec.type == "year":
        year = int(parse_number(raw))
        if not 1950 <= year <= 2100:
            raise ParseError("BAD_YEAR")
        return year
    if spec.type == "date":
        return parse_date(raw)
    if spec.type == "phone":
        try:
            return normalize_phone(to_western_digits(raw), country)
        except AppError as exc:
            raise ParseError("BAD_PHONE") from exc
    return raw.strip()


def _check_sheet(conn: Connection, index: int, sheet: dict[str, Any], country: str, go_live: date) -> _Checked:
    kind: ImportKind = sheet["kind"]
    specs = {spec.key: spec for spec in FIELDS[kind]}
    mapped = {key for key in sheet["column_map"] if key}
    errors: list[RowResult] = []
    missing = [key for key, spec in specs.items() if spec.required and key not in mapped]
    if missing:
        errors.append(RowResult(row_no=0, errors=[RowError(field=key, code="NOT_MAPPED") for key in missing]))
        return _Checked(index, kind, [], errors)
    records = []
    for row_no, cells in _rows(sheet):
        row_errors: list[RowError] = []
        record = _Record(row_no, {})
        for key, spec in specs.items():
            found = cells.get(key, [])
            if not found:
                if spec.required:
                    row_errors.append(RowError(field=key, code="REQUIRED"))
                continue
            for header, raw in found if spec.multiple else found[:1]:
                try:
                    value = _parse(spec, raw, country)
                except ParseError as exc:
                    row_errors.append(RowError(field=key, code=exc.code))
                    continue
                if spec.multiple:
                    if value != 0:
                        record.expenses.append((header, value))
                else:
                    record.values[key] = value
        row_errors += _row_rules(kind, record, go_live)
        if row_errors:
            errors.append(RowResult(row_no=row_no, errors=row_errors))
        else:
            records.append(record)
    errors += _sheet_rules(conn, kind, records)
    bad_rows = {e.row_no for e in errors}
    return _Checked(index, kind, [r for r in records if r.row_no not in bad_rows], errors)


def _row_rules(kind: ImportKind, record: _Record, go_live: date) -> list[RowError]:
    v = record.values
    found: list[RowError] = []
    if kind == "VEHICLES":
        if v.get("purchase_price") is not None and v["purchase_price"] <= 0:
            found.append(RowError(field="purchase_price", code="NEGATIVE"))
        if any(amount < 0 for _, amount in record.expenses):
            found.append(RowError(field="expense", code="NEGATIVE"))
        owed = v.get("owed_to_seller") or ZERO
        if owed < 0 or (v.get("purchase_price") and owed > v["purchase_price"]):
            found.append(RowError(field="owed_to_seller", code="OWED_EXCEEDS_PRICE"))
        if owed > 0 and not v.get("seller_name"):
            found.append(RowError(field="seller_name", code="REQUIRED"))
        if v.get("purchase_date") and v["purchase_date"] > go_live:
            found.append(RowError(field="purchase_date", code="AFTER_GO_LIVE"))
        if v.get("vin"):
            normalized = "".join(ch for ch in v["vin"].upper() if ch.isalnum())
            if not 4 <= len(normalized) <= 20:
                found.append(RowError(field="vin", code="VIN_INVALID"))
            v["vin_key"] = normalized
    elif kind == "INSTALLMENTS":
        paid = v.get("paid") or ZERO
        if v.get("amount") is not None and (v["amount"] <= 0 or paid < 0 or paid > v["amount"]):
            found.append(RowError(field="paid", code="PAID_EXCEEDS_AMOUNT"))
    elif kind == "CASH":
        kind_text = normalize_header(v.get("account_type", ""))
        v["cash_kind"] = "BANK" if any(w in kind_text for w in ("بنك", "bank")) else "CASH_BOX"
    elif kind == "PARTNERS":
        for key in ("capital", "loan_to_partner", "loan_from_partner"):
            if (v.get(key) or ZERO) < 0:
                found.append(RowError(field=key, code="NEGATIVE"))
    return found


def _sheet_rules(conn: Connection, kind: ImportKind, records: list[_Record]) -> list[RowResult]:
    errors: list[RowResult] = []

    def flag(record: _Record, field_name: str, code: str) -> None:
        errors.append(RowResult(row_no=record.row_no, errors=[RowError(field=field_name, code=code)]))

    if kind == "VEHICLES":
        seen: dict[str, int] = {}
        existing = set(
            conn.execute(
                text(
                    "select vin_normalized from public.vehicles where vin_normalized is not null "
                    "and archived_at is null and status not in ('DELIVERED', 'RETURNED_TO_OWNER')"
                )
            ).scalars()
        )
        for record in records:
            key = record.values.get("vin_key")
            if not key:
                continue
            if key in seen:
                flag(record, "vin", "DUPLICATE_IN_FILE")
            elif key in existing:
                flag(record, "vin", "VIN_EXISTS")
            seen.setdefault(key, record.row_no)
    elif kind == "CUSTOMERS":
        phones: set[str] = set()
        existing = set(
            conn.execute(text("select phone_primary from public.customers where phone_primary is not null")).scalars()
        )
        for record in records:
            phone = record.values.get("phone")
            if not phone:
                continue
            if phone in phones:
                flag(record, "phone", "DUPLICATE_IN_FILE")
            elif phone in existing:
                flag(record, "phone", "PHONE_EXISTS")
            phones.add(phone)
    elif kind == "PARTNERS" and records:
        has_shares = conn.execute(text("select exists (select 1 from public.partner_share_history)")).scalar_one()
        if has_shares:
            errors.append(RowResult(row_no=0, errors=[RowError(field=None, code="PARTNERS_EXIST")]))
        total = sum((r.values["percentage"] for r in records), ZERO)
        if total != _HUNDRED:
            errors.append(RowResult(row_no=0, errors=[RowError(field="percentage", code="SHARES_NOT_100")]))
        names = [normalize_header(r.values["name"]) for r in records]
        for record, name in zip(records, names, strict=True):
            if names.count(name) > 1:
                flag(record, "name", "DUPLICATE_IN_FILE")
    return errors


def _check_all(conn: Connection, row: Any) -> list[_Checked]:
    country = customers.country_code(conn)
    return [
        _check_sheet(conn, index, sheet, country, row["go_live_date"])
        for index, sheet in enumerate(row["sheets"])
        if not sheet.get("skip")
    ]


def validate(conn: Connection, job_id: UUID) -> ImportJobOut:
    row = _row(conn, job_id, lock=True)
    if row["status"] == "COMMITTED":
        raise AppError("IMPORT_COMMITTED", "This import has been committed", status_code=409)
    checked = _check_all(conn, row)
    opening = _opening_total(checked)
    validation = {
        "sheets": [
            {"index": c.index, "ok_rows": len(c.records), "errors": [e.model_dump() for e in c.errors]} for c in checked
        ],
        "opening_total": f"{opening:.2f}",
    }
    conn.execute(
        text("update public.import_jobs set status = 'VALIDATED', validation = cast(:v as jsonb) where id = :id"),
        {"v": json.dumps(validation, ensure_ascii=False), "id": job_id},
    )
    return get(conn, job_id)


def _opening_total(checked: list[_Checked]) -> Decimal:
    """Total debits of the opening entry, shown before committing."""
    total = ZERO
    for sheet in checked:
        for record in sheet.records:
            v = record.values
            if sheet.kind == "VEHICLES":
                total += v["purchase_price"] + sum((a for _, a in record.expenses), ZERO)
            elif sheet.kind == "INSTALLMENTS":
                total += v["amount"] - (v.get("paid") or ZERO)
            elif sheet.kind == "CASH":
                total += max(v["balance"], ZERO)
            elif sheet.kind == "PARTNERS":
                total += v.get("loan_to_partner") or ZERO
    return total


# --- Error file --------------------------------------------------------------------------------------


ERROR_LABELS = {
    "REQUIRED": ("مطلوب", "required"),
    "NOT_MAPPED": ("العمود غير محدد", "column not mapped"),
    "BAD_NUMBER": ("رقم غير صحيح", "invalid number"),
    "BAD_DATE": ("تاريخ غير صحيح", "invalid date"),
    "BAD_YEAR": ("سنة غير صحيحة", "invalid year"),
    "BAD_PHONE": ("رقم موبايل غير صحيح", "invalid phone"),
    "BAD_PERCENT": ("نسبة غير صحيحة", "invalid percentage"),
    "NEGATIVE": ("قيمة سالبة", "negative value"),
    "DUPLICATE_IN_FILE": ("مكرر في الملف", "duplicate in the file"),
    "VIN_EXISTS": ("رقم الشاسيه موجود بالفعل", "VIN already in stock"),
    "VIN_INVALID": ("رقم الشاسيه غير صحيح", "invalid VIN"),
    "PHONE_EXISTS": ("الموبايل مسجل لعميل آخر", "phone already used"),
    "OWED_EXCEEDS_PRICE": ("الباقي للبائع أكبر من السعر", "owed exceeds the price"),
    "AFTER_GO_LIVE": ("بعد تاريخ بدء التشغيل", "after the go-live date"),
    "PAID_EXCEEDS_AMOUNT": ("المدفوع أكبر من القسط", "paid exceeds the amount"),
    "PARTNERS_EXIST": ("الشركاء مسجلون بالفعل", "partners already exist"),
    "SHARES_NOT_100": ("مجموع النسب لا يساوي 100%", "shares do not total 100%"),
}


def errors_xlsx(conn: Connection, job_id: UUID, language: str) -> bytes:
    row = _row(conn, job_id)
    checked = _check_all(conn, row)
    side = 0 if language == "ar" else 1
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for result in checked:
        sheet = row["sheets"][result.index]
        out = workbook.create_sheet(title=str(sheet["name"])[:31] or f"Sheet{result.index + 1}")
        out.sheet_view.rightToLeft = language == "ar"
        out.append(["#", *sheet["headers"], "الأخطاء" if side == 0 else "Errors"])
        for cell in out[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E2E8F0")
        by_row = {e.row_no: e.errors for e in result.errors}
        for e in by_row.get(0, []):
            out.append([0, *[None] * len(sheet["headers"]), _describe(e, side)])
        first = sheet["header_row"] + 2
        for offset, values in enumerate(sheet["rows"]):
            row_no = first + offset
            if row_no in by_row:
                out.append([row_no, *values, "؛ ".join(_describe(e, side) for e in by_row[row_no])])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _describe(error: RowError, side: int) -> str:
    label = ERROR_LABELS.get(error.code, (error.code, error.code))[side]
    return f"{error.field}: {label}" if error.field else label


# =====================================================================================================
# Commit (rule 25)
# =====================================================================================================


def _customer_for(conn: Connection, cipher: FieldCipher, name: str, phone: str | None) -> UUID:
    if phone:
        found = conn.execute(
            text("select id from public.customers where phone_primary = :p"), {"p": phone}
        ).scalar_one_or_none()
        if found:
            return found  # type: ignore[no-any-return]
    found = conn.execute(
        text(
            "select id from public.customers where lower(btrim(name)) = lower(btrim(:n)) "
            "and archived_at is null limit 1"
        ),
        {"n": name},
    ).scalar_one_or_none()
    if found:
        return found  # type: ignore[no-any-return]
    return customers.create_customer(conn, CustomerIn(name=name, phone=phone), cipher).id


def commit(conn: Connection, job_id: UUID, cipher: FieldCipher) -> PostingResult[ImportJobOut]:
    row = _row(conn, job_id, lock=True)
    if row["status"] == "COMMITTED":
        raise AppError("IMPORT_COMMITTED", "This import has been committed", status_code=409)
    if row["status"] != "VALIDATED":
        raise AppError("IMPORT_NOT_VALIDATED", "Check the rows before importing", status_code=409)
    info = finance.tenant_info(conn)
    go_live: date = row["go_live_date"]
    finance.check_entry_date(info, go_live)
    checked = _check_all(conn, row)
    error_count = sum(len(c.errors) for c in checked)
    if error_count:
        raise AppError(
            "IMPORT_HAS_ERRORS", "Some rows still have errors", status_code=422, details={"rows": error_count}
        )

    lines: list[Line] = []
    counts = {"vehicles": 0, "customers": 0, "partners": 0, "installment_plans": 0, "cash_accounts": 0}
    purchases: list[tuple[UUID, _Record, UUID | None]] = []
    order = {"CASH": 0, "PARTNERS": 1, "CUSTOMERS": 2, "VEHICLES": 3, "INSTALLMENTS": 4}
    for sheet in sorted(checked, key=lambda c: order[c.kind]):
        if sheet.kind == "CASH":
            for record in sheet.records:
                v = record.values
                account = conn.execute(
                    text(
                        "select id from public.cash_accounts where lower(btrim(name_ar)) = lower(btrim(:n)) "
                        "and archived_at is null"
                    ),
                    {"n": v["account_name"]},
                ).scalar_one_or_none()
                if account is None:
                    account = finance.create_cash_account(
                        conn, CashAccountIn(kind=v["cash_kind"], name_ar=v["account_name"])
                    ).id
                    counts["cash_accounts"] += 1
                ref = finance.active_cash_account(conn, account).ref
                amount = v["balance"]
                lines.append(
                    Line(
                        account=ref.account,
                        debit=amount if amount > 0 else ZERO,
                        credit=-amount if amount < 0 else ZERO,
                        cash_account_id=ref.cash_account_id,
                    )
                )
        elif sheet.kind == "PARTNERS":
            created = []
            for record in sheet.records:
                v = record.values
                partner = partners.create_partner(conn, PartnerIn(name_ar=v["name"], phone=v.get("phone")), cipher)
                created.append((partner.id, v))
                counts["partners"] += 1
                for key, system_key, credit in (
                    ("capital", "PARTNER_CAPITAL", True),
                    ("current_balance", "PARTNER_CURRENT", True),
                    ("loan_from_partner", "PARTNER_LOANS_PAYABLE", True),
                    ("loan_to_partner", "PARTNER_LOANS_RECEIVABLE", False),
                ):
                    amount = v.get(key) or ZERO
                    if amount == 0:
                        continue
                    signed = amount if credit else -amount  # positive = credit
                    lines.append(
                        Line(
                            account=Account.system(system_key),
                            debit=-signed if signed < 0 else ZERO,
                            credit=signed if signed > 0 else ZERO,
                            partner_id=partner.id,
                        )
                    )
            partners.change_shares(
                conn,
                ShareChangeIn.model_validate(
                    {
                        "effective_from": go_live,
                        "shares": [{"partner_id": pid, "percentage": v["percentage"]} for pid, v in created],
                    }
                ),
            )
        elif sheet.kind == "CUSTOMERS":
            for record in sheet.records:
                v = record.values
                customers.create_customer(
                    conn,
                    CustomerIn(
                        name=v["name"],
                        phone=v.get("phone"),
                        national_id=v.get("national_id"),
                        address=v.get("address"),
                        notes=v.get("notes"),
                    ),
                    cipher,
                )
                counts["customers"] += 1
        elif sheet.kind == "VEHICLES":
            for record in sheet.records:
                v = record.values
                vehicle_id = vehicles.create_vehicle(
                    conn,
                    VehicleIn(
                        make=v["make"],
                        model=v["model"],
                        year=v.get("year"),
                        trim=v.get("trim"),
                        color_ext=v.get("color"),
                        vin=v.get("vin"),
                        plate_no=v.get("plate_no"),
                        mileage_km=v.get("mileage_km"),
                        asking_price=v.get("asking_price"),
                        notes=v.get("notes"),
                    ),
                    reason="استيراد أرصدة افتتاحية",
                )
                conn.execute(
                    text("update public.vehicles set stock_date = :d where id = :id"),
                    {"d": v.get("purchase_date") or go_live, "id": vehicle_id},
                )
                lines.append(
                    Line(
                        account=Account.system("VEHICLE_INVENTORY"),
                        debit=v["purchase_price"],
                        vehicle_id=vehicle_id,
                        memo="سعر الشراء",
                    )
                )
                for header, amount in record.expenses:
                    lines.append(
                        Line(
                            account=Account.system("VEHICLE_INVENTORY"),
                            debit=amount,
                            vehicle_id=vehicle_id,
                            memo=header,
                        )
                    )
                seller = None
                owed = v.get("owed_to_seller") or ZERO
                if owed > 0:
                    seller = _customer_for(conn, cipher, v["seller_name"], v.get("seller_phone"))
                    lines.append(
                        Line(
                            account=Account.system("SELLER_PAYABLE"),
                            credit=owed,
                            customer_id=seller,
                            vehicle_id=vehicle_id,
                        )
                    )
                purchases.append((vehicle_id, record, seller))
                counts["vehicles"] += 1
        elif sheet.kind == "INSTALLMENTS":
            groups: dict[tuple[str, str, str], list[_Record]] = {}
            for record in sheet.records:
                v = record.values
                group_key = (
                    normalize_header(v["customer_name"]),
                    v.get("customer_phone") or "",
                    v.get("reference") or "",
                )
                groups.setdefault(group_key, []).append(record)
            for number, items in enumerate(groups.values(), start=1):
                first = items[0].values
                customer_id = _customer_for(conn, cipher, first["customer_name"], first.get("customer_phone"))
                schedule = sorted(
                    ((r.values["due_date"], r.values["amount"] - (r.values.get("paid") or ZERO)) for r in items),
                    key=lambda item: item[0],
                )
                schedule = [(due, remaining) for due, remaining in schedule if remaining > 0]
                if not schedule:
                    continue
                total = sum((remaining for _, remaining in schedule), ZERO)
                plan_id = conn.execute(
                    text(
                        """
                        insert into public.installment_plans
                          (tenant_id, customer_id, financed_amount, frequency, installment_count, first_due_date,
                           opening_reference, opening_vehicle)
                        values (private.current_tenant_id(), :customer, :total, 'MANUAL', :count, :first, :ref,
                                :vehicle)
                        returning id
                        """
                    ),
                    {
                        "customer": customer_id,
                        "total": total,
                        "count": len(schedule),
                        "first": schedule[0][0],
                        "ref": first.get("reference") or f"افتتاحي-{number:03d}",
                        "vehicle": first.get("vehicle"),
                    },
                ).scalar_one()
                for seq, (due, remaining) in enumerate(schedule, start=1):
                    conn.execute(
                        text(
                            "insert into public.installments (tenant_id, plan_id, seq, due_date, amount_due) "
                            "values (private.current_tenant_id(), :plan, :seq, :due, :amount)"
                        ),
                        {"plan": plan_id, "seq": seq, "due": due, "amount": remaining},
                    )
                lines.append(
                    Line(account=Account.system("INSTALLMENT_RECEIVABLE"), debit=total, customer_id=customer_id)
                )
                counts["installment_plans"] += 1

    entry: EntryRef | None = None
    if any(line.debit != line.credit for line in lines):
        draft = rules.opening_balances(
            entry_date=go_live,
            lines=lines,
            description=f"أرصدة افتتاحية — استيراد {row['file_name'] or ''}".strip(),
            source_id=job_id,
        )
        posted = engine.post(conn, draft)
        entry = EntryRef(id=posted.id, entry_no=posted.entry_no)
        for vehicle_id, record, seller in purchases:
            conn.execute(
                text(
                    """
                    insert into public.vehicle_purchases
                      (tenant_id, vehicle_id, seller_customer_id, source, purchase_date, price, deferred_amount,
                       journal_entry_id, notes)
                    values (private.current_tenant_id(), :vehicle, :seller, 'OPENING', :d, :price, :owed, :entry,
                            'رصيد افتتاحي')
                    """
                ),
                {
                    "vehicle": vehicle_id,
                    "seller": seller,
                    "d": record.values.get("purchase_date") or go_live,
                    "price": record.values["purchase_price"],
                    "owed": record.values.get("owed_to_seller") or ZERO,
                    "entry": posted.id,
                },
            )
    for sheet in row["sheets"]:
        if not sheet.get("skip") and not all(h in {f.key for f in FIELDS[sheet["kind"]]} for h in sheet["headers"]):
            _remember(conn, sheet)
    for vehicle_id, _, _ in purchases:
        vehicles.set_status(conn, vehicle_id, "IN_PREPARATION", "استيراد أرصدة افتتاحية")
        vehicles.set_status(conn, vehicle_id, "AVAILABLE", "استيراد أرصدة افتتاحية")
    conn.execute(
        text(
            """
            update public.import_jobs
               set status = 'COMMITTED', committed_at = now(), result = cast(:result as jsonb),
                   journal_entry_id = :entry
             where id = :id
            """
        ),
        {"result": json.dumps(counts), "entry": entry.id if entry else None, "id": job_id},
    )
    return PostingResult[ImportJobOut](document=get(conn, job_id), journal_entries=[entry] if entry else [])


# =====================================================================================================
# Template and opening balance equity (P-08)
# =====================================================================================================


def template_xlsx(language: str) -> bytes:
    """Our template: one sheet per kind, the headers in the language, one example row."""
    examples: dict[ImportKind, dict[str, str]] = {
        "VEHICLES": {
            "make": "Toyota",
            "model": "Corolla",
            "year": "2020",
            "vin": "JTDBR32E720045678",
            "plate_no": "ن ب ل 5678",
            "asking_price": "590000",
            "purchase_date": "2026-09-04",
            "purchase_price": "520000",
            "expense": "15000",
        },
        "CUSTOMERS": {"name": "حسن علي", "phone": "01002223344"},
        "PARTNERS": {"name": "أحمد السيد", "percentage": "50", "capital": "500000"},
        "INSTALLMENTS": {
            "customer_name": "عمر خالد",
            "customer_phone": "01005556677",
            "reference": "C-001",
            "due_date": "2026-10-30",
            "amount": "25000",
        },
        "CASH": {"account_name": "الخزنة الرئيسية", "account_type": "خزنة", "balance": "200000"},
    }
    names = {
        "VEHICLES": ("السيارات", "Vehicles"),
        "CUSTOMERS": ("العملاء", "Customers"),
        "PARTNERS": ("الشركاء", "Partners"),
        "INSTALLMENTS": ("الأقساط", "Installments"),
        "CASH": ("الخزنة والبنوك", "Cash and banks"),
    }
    side = 0 if language == "ar" else 1
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for kind, specs in FIELDS.items():
        sheet = workbook.create_sheet(names[kind][side])
        sheet.sheet_view.rightToLeft = language == "ar"
        sheet.append([spec.label_ar if side == 0 else spec.label_en for spec in specs])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E2E8F0")
        sheet.append([examples[kind].get(spec.key) for spec in specs])
        for index in range(1, len(specs) + 1):
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = 18
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def opening_equity(conn: Connection) -> OpeningEquityOut:
    balance = -Decimal(
        conn.execute(
            text(
                "select coalesce(sum(l.debit - l.credit), 0) from public.journal_lines l "
                "join public.ledger_accounts a on a.id = l.ledger_account_id "
                "where a.system_key = 'OPENING_BALANCE_EQUITY'"
            )
        ).scalar_one()
    )
    info = finance.tenant_info(conn)
    shares = partners.shares_on(conn, info.today)
    return OpeningEquityOut(
        balance=balance,
        partners=[
            {
                "partner_id": str(s.partner_id),
                "name_ar": s.partner_name_ar,
                "name_en": s.partner_name_en,
                "percentage": f"{s.percentage}",
            }
            for s in shares
        ],
    )


def clear_opening_equity(conn: Connection, payload: EquityClearingIn) -> PostingResult[OpeningEquityOut]:
    info = finance.tenant_info(conn)
    finance.check_entry_date(info, payload.clearing_date)
    current = opening_equity(conn)
    total = sum((line.amount for line in payload.lines), ZERO)
    if total <= 0 or total > current.balance:
        raise AppError(
            "OPENING_EQUITY_EXCEEDED",
            "The amounts must not exceed the opening balance equity",
            status_code=422,
            details={"balance": f"{current.balance:.2f}"},
        )
    for line in payload.lines:
        partners.get_partner(conn, line.partner_id)
    draft = rules.opening_equity_clearing(
        entry_date=payload.clearing_date,
        allocations=[(line.partner_id, line.account, line.amount) for line in payload.lines],
        description="توزيع حقوق الملكية الافتتاحية على الشركاء (باتفاقهم)",
        source_id=uuid4(),
    )
    posted = engine.post(conn, draft)
    return PostingResult[OpeningEquityOut](
        document=opening_equity(conn), journal_entries=[EntryRef(id=posted.id, entry_no=posted.entry_no)]
    )
